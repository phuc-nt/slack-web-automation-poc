"""Writes the second batch of Word, Excel and plain-text sample documents.

Usage: python3 kb/build_office_more.py <output directory>
"""

import csv

from openpyxl import Workbook

import content as c
from build_office import OUT, bullets, new_doc, save_book, save_doc, sheet, table

co = c.COMPANY
PLAIN = OUT / "khac"


# --- Word -----------------------------------------------------------------

def quick_starts():
    for p in c.PRODUCTS:
        doc = new_doc(f"Bắt đầu nhanh với {p['name']}", f"{p['kind']} — mã sản phẩm {p['code']}")
        doc.add_paragraph(p["pitch"])
        doc.add_heading("1. Trong hộp", 1)
        bullets(doc, p["box"])
        doc.add_heading("2. Cài đặt trong 4 bước", 1)
        bullets(doc, c.QUICK_START[p["code"]], "List Number")
        doc.add_heading("3. Bảo quản", 1)
        bullets(doc, c.CARE[p["code"]])
        doc.add_heading("4. Khi cần đặt lại", 1)
        doc.add_paragraph(p["reset"])
        doc.add_heading("5. Lỗi hay gặp", 1)
        table(doc, ["Mã lỗi", "Ý nghĩa", "Cách xử lý"], [(code, meaning, action) for code, pc, meaning, action, _ in c.ERROR_CODES if pc == p["code"]])
        doc.add_paragraph(f"Bảo hành {p['warranty']} tháng. Hỗ trợ: {co['hotline']} ({co['hotline_hours']}).")
        save_doc(doc, f"huong-dan/bat-dau-nhanh/Bat-dau-nhanh-{p['name'].replace(' ', '-')}.docx")


def quick_start_en():
    doc = new_doc("Quick Start Guide", "All Moc Lan Home products, English edition")
    names = {"H2": "Plug in power and the network cable, wait for the slow green blink, then scan the QR code under the hub in the app.",
             "C5": "Power on, wait for the camera to rotate once, then show the QR code from the app to the lens at 15–20 cm. 2.4 GHz Wi-Fi only.",
             "O7": "Test indoors before drilling. Power with 12 V or PoE and mount at 2.5–3 m using the waterproof junction box.",
             "L9": "Insert four AA alkaline batteries, change the default admin code 123456, then enrol fingerprints (three touches each).",
             "M3": "Turn off the modem's Wi-Fi, plug the main unit into the modem's LAN port and create the network in the app.",
             "P1": "Plug in, turn on Bluetooth on the phone, add the plug in the app and choose a 2.4 GHz network. Maximum load 2,500 W.",
             "S4": "Pull the battery tab, hold the pairing button for 5 seconds and add it through Hub H2. Battery type CR2450.",
             "D6": "Charge fully (about 5 hours), plug in the indoor chime within 10 m of the door and pair with the QR code."}
    for p in c.PRODUCTS:
        doc.add_heading(p["name"], 1)
        doc.add_paragraph(names[p["code"]])
        doc.add_paragraph(f"Warranty: {p['warranty']} months.")
    doc.add_heading("Support", 1)
    doc.add_paragraph(f"Hotline {co['hotline']} or {co['email']}. English-speaking agents: 9:00–18:00, Monday to Friday (GMT+7).")
    save_doc(doc, "huong-dan/Quick-Start-Guide-EN.docx")


def loyalty():
    doc = new_doc(f"Chương trình hội viên {c.LOYALTY_NAME}", "Điều khoản áp dụng từ 01/07/2026")
    doc.add_heading("1. Hạng hội viên", 1)
    table(doc, ["Hạng", "Chi tiêu trong 12 tháng", "Chiết khấu", "Quyền lợi"], [(t, "Từ " + c.vnd(s) if s else "Đăng ký tài khoản", d, perk) for t, s, d, perk in c.LOYALTY])
    doc.add_heading("2. Quy định", 1)
    bullets(doc, c.LOYALTY_RULES)
    doc.add_heading("3. Phân biệt với gói Cloud", 1)
    doc.add_paragraph(f"{c.LOYALTY_NAME} tính theo tiền mua thiết bị. Gói Cloud Family và Pro là dịch vụ trả phí hằng tháng; "
                      f"quyền đổi trả {c.RETURNS['member_days']} ngày gắn với gói Cloud, không gắn với hạng {c.LOYALTY_NAME}.")
    save_doc(doc, "chinh-sach/Chuong-trinh-hoi-vien-Moc-Lan-Club.docx")


def dealer_policy():
    doc = new_doc("Chính sách đại lý 2026", "Dành cho đại lý Cấp 1 và Cấp 2")
    doc.add_heading("1. Chiết khấu và cam kết", 1)
    table(doc, ["Cấp", "Chiết khấu", "Doanh số cam kết", "Thời hạn thanh toán", "Ưu tiên"], c.DEALER_TERMS)
    doc.add_heading("2. Quy định chung", 1)
    bullets(doc, c.DEALER_RULES)
    doc.add_heading("3. Lịch đào tạo", 1)
    table(doc, ["Ngày", "Nơi tổ chức", "Nội dung", "Số chỗ"], c.DEALER_TRAINING)
    doc.add_heading("4. Liên hệ", 1)
    doc.add_paragraph("Phòng Kinh doanh đại lý: daily@moclanhome.example. Đơn đặt hàng gửi trước 15:00 được xuất kho trong ngày.")
    save_doc(doc, "dich-vu/Chinh-sach-dai-ly-2026.docx")


def privacy():
    doc = new_doc("Chính sách bảo mật dữ liệu cá nhân", "Bản dành cho khách hàng, cập nhật 01/09/2026")
    for i, line in enumerate(c.PRIVACY, 1):
        doc.add_paragraph(f"{i}. {line}")
    doc.add_heading("Quyền của khách hàng", 1)
    bullets(doc, ["Xem và tải về dữ liệu của mình.", "Yêu cầu sửa thông tin sai.", "Yêu cầu xoá tài khoản.", "Rút lại đồng ý nhận tin quảng cáo bất cứ lúc nào."])
    save_doc(doc, "chinh-sach/Chinh-sach-bao-mat-du-lieu.docx")

    en = new_doc("Privacy Notice", "English summary, updated 1 September 2026")
    bullets(en, ["We collect your phone number, email, delivery address, device serial numbers and device activity logs.",
                 "Video is stored only when you insert a memory card or subscribe to a Cloud plan.", "We do not sell personal data.",
                 "Device logs are kept for 90 days; invoices for 10 years as accounting law requires.",
                 "You can delete your account in the app; the account and all video are erased within 30 days.",
                 "Data requests: baomat@moclanhome.example, answered within 72 hours."])
    save_doc(en, "chinh-sach/Privacy-Notice-EN.docx")


def infosec():
    doc = new_doc("Quy định an toàn thông tin cho tổng đài", "Tài liệu nội bộ, bắt buộc đọc trong tuần đầu")
    bullets(doc, c.INFOSEC_RULES, "List Number")
    doc.add_heading("Vi phạm", 1)
    doc.add_paragraph("Vi phạm lần đầu ở mức nhẹ bị nhắc nhở bằng văn bản. Chép dữ liệu khách hàng ra ngoài là vi phạm nghiêm trọng và bị xử lý theo nội quy lao động.")
    save_doc(doc, "noi-bo/Quy-dinh-an-toan-thong-tin.docx")


def incidents():
    for code, date, title, duration, impact, cause, fix, comp in c.INCIDENTS:
        doc = new_doc(f"Báo cáo sự cố {code}", f"{title} — ngày {date}")
        table(doc, ["Mục", "Nội dung"], [("Mã sự cố", code), ("Ngày", date), ("Thời gian ảnh hưởng", duration), ("Phạm vi ảnh hưởng", impact),
                                          ("Nguyên nhân gốc", cause), ("Cách khắc phục", fix), ("Bồi thường", comp)])
        doc.add_heading("Hướng dẫn cho tổng đài", 1)
        bullets(doc, ["Dùng mẫu trả lời SC-01 khi sự cố đang diễn ra.", "Không suy đoán nguyên nhân với khách khi chưa có thông báo chính thức.",
                      f"Khách hỏi về bồi thường: {comp.lower()}."])
        save_doc(doc, f"noi-bo/su-co/Bao-cao-su-co-{code}.docx")


def scenes_guide():
    doc = new_doc("Sáu kịch bản tự động mẫu", "Gợi ý cài đặt cho khách mới dùng hub")
    for name, trigger, actions, devices in c.SCENES:
        doc.add_heading(name, 2)
        table(doc, ["Mục", "Nội dung"], [("Khi", trigger), ("Thì", actions), ("Thiết bị cần có", devices)])
    doc.add_paragraph("Kịch bản được tạo trong ứng dụng, mục Tự động > Thêm kịch bản. Mỗi hub lưu tối đa 200 kịch bản.")
    save_doc(doc, "huong-dan/Kich-ban-tu-dong-mau.docx")


def app_howto():
    doc = new_doc("Hướng dẫn thao tác trên ứng dụng", "Tám việc khách hay hỏi nhất")
    for title, steps, note in c.HOWTO:
        doc.add_heading(title, 1)
        bullets(doc, steps, "List Number")
        doc.add_paragraph(f"Lưu ý: {note}")
    save_doc(doc, "huong-dan/Thao-tac-tren-ung-dung.docx")


def email_templates():
    doc = new_doc("Mẫu trả lời nhanh", "Dùng cho email, chat và kết thúc cuộc gọi")
    table(doc, ["Mã mẫu", "Tình huống", "Nội dung"], c.MACROS)
    doc.add_paragraph("Phần trong ngoặc vuông phải được thay bằng thông tin thật trước khi gửi.")
    save_doc(doc, "noi-bo/Mau-tra-loi-nhanh.docx")


def work_rules():
    doc = new_doc("Nội quy làm việc của bộ phận CSKH", "Áp dụng cho nhân viên chính thức và thử việc")
    doc.add_heading("1. Giờ làm việc", 1)
    table(doc, ["Ca", "Giờ", "Phụ cấp"], [("Sáng", "8:00–15:00", "Không"), ("Chiều", "14:00–21:00", "Không"), ("Chat tối", "15:00–22:00", "40.000 đ mỗi ca"),
                                          ("Lễ, Tết", "Theo lịch riêng", "300% lương ngày")])
    doc.add_heading("2. Nghỉ phép", 1)
    bullets(doc, ["Nghỉ phép năm 14 ngày, đăng ký trước 7 ngày.", "Tháng 1 và tháng 2 hạn chế nghỉ vì cao điểm Tết.", "Nghỉ ốm báo trưởng nhóm trước giờ vào ca 2 giờ.",
                  "Đổi ca báo trước 48 giờ, tối đa 3 ca mỗi tháng."])
    doc.add_heading("3. Thử việc và đánh giá", 1)
    bullets(doc, ["Thử việc 2 tháng, hưởng 85% lương.", "Đánh giá chất lượng: mỗi tuần chấm 5 cuộc gọi ngẫu nhiên, điểm đạt là 85 trên 100.",
                  "Thưởng quý cho nhóm đạt mức hài lòng từ 94% trở lên: 1.500.000 đ mỗi người."])
    doc.add_heading("4. Làm việc từ xa", 1)
    doc.add_paragraph("Nhân viên qua thử việc được làm từ xa tối đa 2 ngày mỗi tuần với kênh chat và email; kênh thoại chỉ làm tại văn phòng.")
    save_doc(doc, "noi-bo/Noi-quy-lam-viec-CSKH.docx")


# --- Excel ----------------------------------------------------------------

def monthly_stats():
    reasons = [("Camera ngoại tuyến (E201)", 0.18), ("Gói Cloud và hoá đơn", 0.15), ("Khóa báo pin yếu (E401)", 0.11), ("Tiến độ bảo hành", 0.09),
               ("Ghép thiết bị với hub", 0.08), ("Đổi trả", 0.07), ("Giao hàng", 0.06), ("Khác", 0.26)]
    for month, total, phone, chat, email, csat, fcr, top in c.MONTHLY:
        wb = Workbook()
        ws = sheet(wb, "Tổng hợp", ["Chỉ số", "Giá trị"], [("Tháng", month), ("Tổng liên hệ", total), ("Điện thoại", phone), ("Chat", chat), ("Email", email),
                                                          ("Mức hài lòng (CSAT)", csat), ("Giải quyết ngay lần đầu", fcr), ("Lý do nhiều nhất", top)], first=True)
        ws.append(["Kiểm tra tổng", "=B4+B5+B6"])
        sheet(wb, "Theo lý do", ["Lý do", "Tỷ lệ", "Số liên hệ (ước tính)"], [(name, f"{share:.0%}", round(total * share)) for name, share in reasons])
        sheet(wb, "Theo kênh", ["Kênh", "Số liên hệ", "Tỷ trọng"], [(k, v, f"{v / total:.1%}") for k, v in (("Điện thoại", phone), ("Chat", chat), ("Email", email))])
        save_book(wb, f"bao-cao/Thong-ke-lien-he-{month[3:]}-{month[:2]}.xlsx")


def weekly_kpi():
    wb = Workbook()
    rows = [(f"Tuần {27 + i}", 3_420 + i * 46, f"{91.6 + i * 0.2:.1f}%".replace(".", ","), f"{77.8 + i * 0.25:.1f}%".replace(".", ","), f"{92 - i} giây") for i in range(13)]
    sheet(wb, "KPI theo tuần", ["Tuần (2026)", "Số liên hệ", "CSAT", "Giải quyết lần đầu", "Thời gian chờ dài nhất"], rows, first=True)
    sheet(wb, "Mục tiêu", ["Chỉ số", "Mục tiêu"], [("CSAT", c.KPI["csat_target"]), ("Giải quyết lần đầu", c.KPI["first_call_resolution"]), ("Nhấc máy", c.KPI["answer_within"]),
                                                   ("Giữ máy tối đa", c.KPI["max_hold"]), ("Gọi lại", c.KPI["callback"])])
    save_book(wb, "bao-cao/KPI-theo-tuan-Q3-2026.xlsx")


def stock():
    wb = Workbook()
    ws = sheet(wb, "Tồn kho 01-10-2026", ["Mã", "Sản phẩm", *[w for w, _ in c.WAREHOUSES], "Tổng"],
               [(code, c.product(code)["name"], *qty, None) for code, qty in c.STOCK.items()], first=True)
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=6, value=f"=SUM(C{r}:E{r})")
    sheet(wb, "Kho", ["Kho", "Địa chỉ"], c.WAREHOUSES)
    sheet(wb, "Quy tắc", ["Mục", "Nội dung"], [("Ngưỡng báo sắp hết", "Dưới 100 chiếc ở một kho"), ("Chuyển kho", "Mất 2 ngày giữa Kho Bắc và Kho Nam, 1 ngày tới Kho Trung"),
                                               ("Kiểm kê", "Ngày 1 hằng tháng"), ("Hàng đặt trước", "Giữ tối đa 72 giờ chờ thanh toán")])
    save_book(wb, "bang-gia/Ton-kho-theo-kho.xlsx")


def roster():
    names = ["Thu An", "Bảo Châu", "Minh Đức", "Hải Yến", "Quốc Huy", "Ngọc Lan", "Tuấn Kiệt", "Mai Phương", "Gia Bảo", "Khánh Linh", "Đức Anh", "Thanh Tâm"]
    shifts = ["Sáng", "Chiều", "Chat tối", "Nghỉ"]
    wb = Workbook()
    days = [f"{d:02d}/11" for d in range(1, 15)]
    sheet(wb, "Lịch trực 01-14.11", ["Nhân viên", "Số nội bộ", *days],
          [(name, 5001 + i, *[shifts[(i + d) % 4] for d in range(14)]) for i, name in enumerate(names)], first=True)
    sheet(wb, "Trưởng ca", ["Ca", "Trưởng ca", "Số nội bộ"], [("Sáng", "Thu An", 5001), ("Chiều", "Quốc Huy", 5005), ("Chat tối", "Gia Bảo", 5009)])
    sheet(wb, "Ghi chú", ["Mục", "Nội dung"], [("Ngày 11/11", "Tăng cường 4 người ca chiều vì chương trình Ngày đôi 11.11"), ("Kỹ thuật viên trực Lock L9", "Số nội bộ 8009, trực 24/7")])
    save_book(wb, "noi-bo/Lich-truc-thang-11-2026.xlsx")


def quiz():
    wb = Workbook()
    sheet(wb, "Câu hỏi", ["STT", "Câu hỏi", "Đáp án"], [(i, q, a) for i, (q, a) in enumerate(c.QUIZ, 1)], first=True)
    sheet(wb, "Quy định", ["Mục", "Nội dung"], [("Số câu trong bài thi", "40 câu, rút ngẫu nhiên từ ngân hàng câu hỏi"), ("Điểm đạt", "Từ 32 câu đúng"),
                                                ("Thời gian", "45 phút"), ("Thi lại", "Sau 3 ngày, tối đa 2 lần")])
    save_book(wb, "noi-bo/Ngan-hang-cau-hoi-dao-tao.xlsx")


def spec_compare():
    keys = []
    for p in c.PRODUCTS:
        keys += [k for k, _ in p["specs"] if k not in keys]
    wb = Workbook()
    sheet(wb, "So sánh thông số", ["Thông số", *[p["name"] for p in c.PRODUCTS]],
          [(k, *[dict(p["specs"]).get(k, "—") for p in c.PRODUCTS]) for k in keys], first=True)
    sheet(wb, "Giá và bảo hành", ["Sản phẩm", "Giá niêm yết (đ)", "Bảo hành (tháng)", "Khôi phục cài đặt gốc"], [(p["name"], p["price"], p["warranty"], p["reset"]) for p in c.PRODUCTS])
    save_book(wb, "ky-thuat/So-sanh-thong-so.xlsx")


def loyalty_book():
    wb = Workbook()
    sheet(wb, "Hạng hội viên", ["Hạng", "Chi tiêu 12 tháng từ (đ)", "Chiết khấu", "Quyền lợi"], c.LOYALTY, first=True)
    sheet(wb, "Quy định", ["STT", "Quy định"], list(enumerate(c.LOYALTY_RULES, 1)))
    save_book(wb, "chinh-sach/Hang-hoi-vien.xlsx")


def scenes_book():
    wb = Workbook()
    sheet(wb, "Kịch bản mẫu", ["Tên", "Khi", "Thì", "Thiết bị cần có"], c.SCENES, first=True)
    save_book(wb, "ky-thuat/Kich-ban-tu-dong.xlsx")


def dealer_book():
    wb = Workbook()
    sheet(wb, "Điều khoản", ["Cấp", "Chiết khấu", "Doanh số cam kết", "Thời hạn thanh toán", "Ưu tiên"], c.DEALER_TERMS, first=True)
    sheet(wb, "Giá nhập", ["Sản phẩm", "Giá niêm yết (đ)", "Cấp 1 (đ)", "Cấp 2 (đ)"],
          [(p["name"], p["price"], round(p["price"] * 0.78, -3), round(p["price"] * 0.85, -3)) for p in c.PRODUCTS])
    sheet(wb, "Lịch đào tạo", ["Ngày", "Nơi tổ chức", "Nội dung", "Số chỗ"], c.DEALER_TRAINING)
    sheet(wb, "Danh sách đại lý", ["Mã đại lý", "Tên đại lý", "Tỉnh thành", "Cấp"], c.DEALERS)
    save_book(wb, "dich-vu/Bang-gia-va-dieu-khoan-dai-ly.xlsx")


def centre_book():
    wb = Workbook()
    sheet(wb, "Chi tiết trung tâm", ["Thành phố", "Số nội bộ quản lý", "Gửi xe", "Số quầy", "Tủ nhận thiết bị", "Ghi chú"],
          [(city, *detail) for city, detail in c.CENTRE_DETAILS.items()], first=True)
    save_book(wb, "dich-vu/Chi-tiet-trung-tam.xlsx")


# --- Plain formats --------------------------------------------------------

def plain_files():
    PLAIN.mkdir(parents=True, exist_ok=True)
    with open(PLAIN / "Ma-loi.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["ma_loi", "san_pham", "y_nghia", "huong_dan", "cap_xu_ly"])
        writer.writerows((code, c.product(pc)["name"], meaning, action, level) for code, pc, meaning, action, level in c.ERROR_CODES)
    with open(PLAIN / "Ton-kho.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["ma", "san_pham", "kho_bac", "kho_trung", "kho_nam"])
        writer.writerows((code, c.product(code)["name"], *qty) for code, qty in c.STOCK.items())
    with open(PLAIN / "Cau-hoi-dao-tao.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["cau_hoi", "dap_an"])
        writer.writerows(c.QUIZ)

    z, g = c.ZALO, c.GREEN_BATTERY
    (PLAIN / "Kenh-ho-tro-Zalo-OA.md").write_text(
        "# Kênh hỗ trợ Zalo OA\n\n"
        f"- Bắt đầu: {z['start']}\n- Giờ trực: {z['hours']}\n- Thời gian trả lời: {z['reply']}\n- Phạm vi: {z['scope']}\n\n"
        "## Lưu ý cho tổng đài viên\n\nKhông gửi mã OTP, mật khẩu hay đường dẫn thanh toán qua Zalo. Mỗi hội thoại vẫn phải có phiếu CS-2026.\n", encoding="utf-8")
    (PLAIN / "Chuong-trinh-Pin-Xanh.md").write_text(
        "# Chương trình Pin Xanh\n\n"
        f"Thu hồi pin cũ tại cả 4 trung tâm bảo hành từ tháng {g['start']}.\n\n- {g['reward']}.\n- {g['limit']}.\n"
        "- Nhận mọi loại pin AA, AAA, CR2450 và pin sạc của Bell D6; không nhận pin phồng, pin phồng phải báo Cấp 2.\n", encoding="utf-8")
    (PLAIN / "Cau-hoi-thuong-gap-chung.md").write_text(
        "# Câu hỏi thường gặp chung\n\n" + "".join(f"## {q}\n\n{a}\n\n" for q, a in c.QUIZ[:12]), encoding="utf-8")
    (PLAIN / "Thong-bao-nghi-Tet-2027.txt").write_text(
        "THÔNG BÁO LỊCH NGHỈ TẾT 2027\n\n"
        "Trung tâm bảo hành nghỉ từ 28 tháng Chạp đến hết mùng 4 Tết.\n"
        f"Tổng đài {co['hotline']} vẫn trực 8:00–21:00 suốt kỳ nghỉ.\n"
        "Đơn bảo hành nhận sau ngày 20/01/2027 được hẹn trả sau Tết.\n"
        "Đơn hàng đặt từ 25 tháng Chạp được giao từ mùng 5 Tết.\n", encoding="utf-8")
    (PLAIN / "Thong-bao-ngung-linh-kien.txt").write_text(
        "THÔNG BÁO NGỪNG CẤP LINH KIỆN\n\n"
        "Từ ngày 31/12/2026, Mộc Lan Home ngừng cấp linh kiện cho Cam C3 và Hub H1.\n"
        "Khách đang dùng hai sản phẩm này được giảm 25% khi đổi sang Cam C5 hoặc Hub H2 đến hết 31/03/2027, mã LENDOI25.\n"
        "Thiết bị cũ được thu hồi miễn phí tại 4 trung tâm bảo hành.\n", encoding="utf-8")
    rows = "".join(f"<tr><td>{name}</td><td>{trigger}</td><td>{actions}</td><td>{devices}</td></tr>" for name, trigger, actions, devices in c.SCENES)
    (PLAIN / "Kich-ban-tu-dong.html").write_text(
        "<!doctype html><html lang=\"vi\"><head><meta charset=\"utf-8\"><title>Kịch bản tự động mẫu</title></head><body><h1>Kịch bản tự động mẫu</h1>"
        f"<table><tr><th>Tên</th><th>Khi</th><th>Thì</th><th>Thiết bị</th></tr>{rows}</table></body></html>", encoding="utf-8")
    tiers = "".join(f"<li><strong>{t}</strong>: chiết khấu {d}. {perk}.</li>" for t, _, d, perk in c.LOYALTY)
    (PLAIN / "Hoi-vien-Moc-Lan-Club.html").write_text(
        f"<!doctype html><html lang=\"vi\"><head><meta charset=\"utf-8\"><title>{c.LOYALTY_NAME}</title></head><body><h1>{c.LOYALTY_NAME}</h1>"
        f"<ul>{tiers}</ul><h2>Quy định</h2><ul>{''.join(f'<li>{r}</li>' for r in c.LOYALTY_RULES)}</ul></body></html>", encoding="utf-8")


def main():
    for build in (quick_starts, quick_start_en, loyalty, dealer_policy, privacy, infosec, incidents, scenes_guide, app_howto, email_templates, work_rules,
                  monthly_stats, weekly_kpi, stock, roster, quiz, spec_compare, loyalty_book, scenes_book, dealer_book, centre_book, plain_files):
        build()
    print(f"{sum(1 for p in OUT.rglob('*') if p.is_file())} files in {OUT}")


if __name__ == "__main__":
    main()
