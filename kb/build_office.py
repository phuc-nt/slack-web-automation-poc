"""Writes the Word, Excel and plain-text sample documents.

Usage: python3 kb/build_office.py <output directory>
Needs python-docx and openpyxl.
"""

import csv
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import content as c

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "artifacts/kb/sharepoint")


# --- Word -----------------------------------------------------------------

def new_doc(title, subtitle):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    doc.core_properties.title = title
    doc.core_properties.author = c.COMPANY["name"]
    doc.add_heading(title, 0)
    doc.add_paragraph(subtitle).italic = True
    return doc


def bullets(doc, items, style="List Bullet"):
    for item in items:
        doc.add_paragraph(item, style=style)


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    doc.add_paragraph()


def save_doc(doc, path):
    target = OUT / path
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def warranty_policy():
    w = c.WARRANTY
    doc = new_doc("Chính sách bảo hành", f"Áp dụng cho sản phẩm {c.COMPANY['name']} mua từ ngày 01/01/2026")
    doc.add_heading("1. Thời hạn bảo hành theo sản phẩm", 1)
    table(doc, ["Sản phẩm", "Loại", "Thời hạn"], [(p["name"], p["kind"], f"{p['warranty']} tháng") for p in c.PRODUCTS])
    doc.add_paragraph("Thời hạn tính từ ngày giao hàng ghi trên hoá đơn. Phụ kiện có thời hạn riêng, xem bảng giá phụ kiện.")
    doc.add_heading("2. Quyền lợi", 1)
    bullets(doc, [w["one_for_one"], f"Thời gian xử lý: {w['process_days']}.", w["loaner"], w["extended"]])
    doc.add_heading("3. Trường hợp không được bảo hành", 1)
    bullets(doc, w["exclusions"])
    doc.add_heading("4. Dịch vụ ngoài bảo hành", 1)
    doc.add_paragraph(f"Phí kiểm tra thiết bị hết bảo hành là {c.vnd(w['out_of_warranty_check'])}, được trừ vào phí sửa nếu khách đồng ý sửa. "
                      f"Thay bộ pin và gioăng cho Lock L9 có giá {c.vnd(w['battery_swap_l9'])}.")
    doc.add_heading("5. Cách gửi bảo hành", 1)
    bullets(doc, ["Mở ứng dụng, vào Hỗ trợ > Tạo yêu cầu bảo hành và nhập số sê-ri.",
                  "Nhận mã yêu cầu dạng BH-2026-xxxxxx qua tin nhắn.",
                  "Mang thiết bị đến trung tâm hoặc gửi bưu điện; phí gửi hai chiều do Mộc Lan Home trả khi còn bảo hành.",
                  f"Theo dõi tiến độ trong ứng dụng hoặc gọi {c.COMPANY['hotline']}."], "List Number")
    save_doc(doc, "chinh-sach/Chinh-sach-bao-hanh.docx")


def cloud_terms():
    doc = new_doc("Điều khoản dịch vụ Cloud", "Gói lưu trữ video và tính năng thông minh")
    doc.add_heading("1. Các gói dịch vụ", 1)
    table(doc, ["Gói", "Giá mỗi tháng", "Lưu video", "Số camera", "Thời gian phản hồi hỗ trợ"],
          [(n, "Miễn phí" if p == 0 else c.vnd(p), d, cams, sla) for n, p, d, cams, sla in c.PLANS])
    doc.add_paragraph(c.PLAN_YEARLY_NOTE)
    doc.add_heading("2. Gia hạn và huỷ", 1)
    bullets(doc, ["Gói tự gia hạn vào ngày đăng ký hằng tháng; thông báo được gửi trước 3 ngày.",
                  "Khách huỷ bất cứ lúc nào trong Tài khoản > Gói Cloud; gói còn hiệu lực đến hết kỳ đã trả.",
                  "Hoàn tiền theo tỷ lệ chỉ áp dụng cho gói năm, trong 14 ngày đầu.",
                  "Khi hạ gói, video cũ hơn thời hạn của gói mới bị xoá sau 72 giờ."])
    doc.add_heading("3. Dữ liệu và quyền riêng tư", 1)
    bullets(doc, ["Video được mã hoá AES-256 khi lưu và TLS 1.3 khi truyền.",
                  "Máy chủ đặt tại Việt Nam. Nhân viên không xem được video nếu khách chưa cấp quyền bằng mã 6 số.",
                  "Khách tải toàn bộ dữ liệu của mình trong Tài khoản > Quyền riêng tư > Xuất dữ liệu; tệp sẵn sàng trong 48 giờ.",
                  "Tài khoản bị xoá theo yêu cầu trong 30 ngày, kèm toàn bộ video."])
    doc.add_heading("4. Giới hạn sử dụng hợp lý", 1)
    doc.add_paragraph("Mỗi camera được tải lên tối đa 60 GB mỗi tháng. Vượt mức này, camera chuyển sang ghi theo sự kiện đến đầu tháng sau.")
    save_doc(doc, "chinh-sach/Dieu-khoan-dich-vu-Cloud.docx")


def returns_en():
    r = c.RETURNS
    doc = new_doc("Returns and Refunds Policy", "English summary for international customers and partners")
    doc.add_heading("Return window", 1)
    doc.add_paragraph(f"You may return a product within {r['days']} days of delivery. Family and Pro subscribers have {r['member_days']} days. "
                      "The refund reaches your original payment method within 5 to 7 working days after our warehouse receives the item.")
    doc.add_heading("Conditions", 1)
    bullets(doc, ["The item is in its original box with all accessories and an intact MLH-SEAL sticker.",
                  "You have the invoice or the order number.", "The item has no scratches, water damage or drill marks."])
    doc.add_heading("Items that cannot be returned", 1)
    bullets(doc, ["Opened memory cards and batteries.", "Clearance items.",
                  "An ML Lock L9 that has been fitted to a door, unless it has a technical fault."])
    doc.add_heading("Shipping cost", 1)
    doc.add_paragraph("Free when the product is faulty. If you changed your mind, the return shipping fee is 35,000 VND. "
                      "International orders are returned at the customer's cost; customs duties are not refunded.")
    doc.add_heading("How to start a return", 1)
    bullets(doc, ["Open the app and go to Orders > Return an item.", "Print the return label (code format DT-2026-xxxxxx).",
                  f"Questions: {c.COMPANY['email']}. English support is available 9:00–18:00, Monday to Friday (GMT+7)."], "List Number")
    save_doc(doc, "chinh-sach/Returns-and-Refunds-Policy-EN.docx")


def lock_install_guide():
    p = c.product("L9")
    doc = new_doc("Hướng dẫn lắp đặt ML Lock L9", "Dành cho kỹ thuật viên và khách tự lắp")
    doc.add_heading("1. Trước khi lắp", 1)
    bullets(doc, ["Cửa dày 38–90 mm, khoảng cách từ mép cửa đến tâm lỗ khoá (backset) 60 hoặc 70 mm.",
                  "Cửa không cong vênh quá 3 mm; cửa xệ phải chỉnh bản lề trước.",
                  "Dụng cụ: khoan, mũi khoét 54 mm, mũi 25 mm, đục, tua vít, thước, mẫu khoan trong hộp."])
    doc.add_heading("2. Trong hộp có gì", 1)
    bullets(doc, p["box"])
    doc.add_heading("3. Các bước lắp", 1)
    bullets(doc, ["Dán mẫu khoan lên cửa ở độ cao 1.000 mm tính từ sàn đến tâm tay nắm.",
                  "Khoét lỗ 54 mm cho thân khoá và lỗ 25 mm ở cạnh cửa cho chốt.",
                  "Lắp ruột khoá, mũi tên trên ruột hướng lên trên.",
                  "Luồn cáp tín hiệu qua lỗ, gắn thân trước, rồi thân sau; không kẹp cáp.",
                  "Siết 4 vít theo đường chéo, lực vừa phải.",
                  "Lắp 4 pin AA kiềm và thử chốt 5 lần khi cửa đang mở."], "List Number")
    doc.add_heading("4. Cài đặt lần đầu", 1)
    bullets(doc, ["Mã quản trị mặc định là 123456; khoá buộc đổi mã ở lần dùng đầu tiên.",
                  "Đăng ký vân tay quản trị: nhấn * rồi #, nhập mã quản trị, chọn 1.",
                  "Mỗi vân tay đặt 3 lần ở 3 góc khác nhau.",
                  "Ghép với ứng dụng bằng Bluetooth; muốn mở từ xa cần Hub H2 hoặc cổng Wi-Fi của khoá."])
    doc.add_heading("5. Lưu ý an toàn", 1)
    bullets(doc, ["Luôn giữ một chìa cơ ở ngoài nhà.", "Không dùng pin sạc: điện áp 1,2 V làm khoá báo pin yếu sai.",
                  f"Khôi phục cài đặt gốc: {p['reset']}", "Thời gian lắp chuẩn là 90 phút với cửa gỗ."])
    save_doc(doc, "huong-dan/Huong-dan-lap-dat-Lock-L9.docx")


def hub_manual():
    p = c.product("H2")
    doc = new_doc("Hướng dẫn sử dụng ML Hub H2", "Phiên bản firmware 3.4.0")
    doc.add_heading("1. Giới thiệu", 1)
    doc.add_paragraph(p["pitch"])
    table(doc, ["Thông số", "Giá trị"], p["specs"])
    doc.add_heading("2. Cài đặt", 1)
    bullets(doc, ["Cắm nguồn và cáp mạng vào router; chờ đèn xanh lá nháy chậm.",
                  "Mở ứng dụng, chọn Thêm thiết bị > Hub, quét mã QR ở đáy hub.",
                  "Đặt tên phòng. Sau khi ghép xong có thể rút cáp mạng và dùng Wi-Fi 2,4 GHz."], "List Number")
    doc.add_heading("3. Kịch bản tự động", 1)
    bullets(doc, ["Mỗi hub lưu tối đa 200 kịch bản; mỗi kịch bản có tối đa 10 điều kiện và 20 hành động.",
                  "Kịch bản dùng thiết bị Zigbee vẫn chạy khi mất Internet. Kịch bản dùng thời tiết hoặc vị trí điện thoại cần Internet.",
                  "Ví dụ: khi Sensor S4 báo mở cửa sau 22:00 thì bật Plug P1 ở hành lang trong 5 phút."])
    doc.add_heading("4. Ý nghĩa đèn báo", 1)
    table(doc, ["Trạng thái đèn", "Ý nghĩa"], c.LED_STATES)
    doc.add_heading("5. Sao lưu và chuyển hub", 1)
    doc.add_paragraph("Hub tự sao lưu cấu hình lên Cloud lúc 3:00 mỗi ngày. Khi đổi sang hub mới, chọn Khôi phục từ bản sao lưu; "
                      "thiết bị Zigbee không cần ghép lại nếu hub cũ đã được gỡ khỏi tài khoản.")
    doc.add_heading("6. Khôi phục cài đặt gốc", 1)
    doc.add_paragraph(p["reset"])
    save_doc(doc, "huong-dan/Huong-dan-su-dung-Hub-H2.docx")


def complaint_process():
    doc = new_doc("Quy trình xử lý khiếu nại", "Tài liệu nội bộ của bộ phận Chăm sóc khách hàng")
    doc.add_heading("1. Phân loại", 1)
    table(doc, ["Mức", "Ví dụ", "Thời hạn phản hồi đầu tiên", "Thời hạn giải quyết"],
          [("Thường", "Giao trễ, hỏi về hoá đơn", "4 giờ làm việc", "3 ngày làm việc"),
           ("Nghiêm trọng", "Thiết bị lỗi lặp lại, hoàn tiền chậm", "1 giờ làm việc", "5 ngày làm việc"),
           ("Khẩn cấp", "Khoá cửa không mở được, nghi lộ dữ liệu", "15 phút", "24 giờ")])
    doc.add_heading("2. Các cấp hỗ trợ", 1)
    table(doc, ["Cấp", "Ai xử lý", "Phạm vi", "Thời hạn chuyển tiếp"], c.SUPPORT_LEVELS)
    doc.add_heading("3. Quyền bồi thường của tổng đài viên", 1)
    bullets(doc, ["Tổng đài viên được tặng mã giảm giá tối đa 100.000 đ mà không cần duyệt.",
                  "Trưởng ca duyệt bồi thường đến 500.000 đ hoặc 1 tháng gói Cloud.",
                  "Trên 500.000 đ cần Trưởng phòng Chăm sóc khách hàng duyệt trong 1 ngày làm việc."])
    doc.add_heading("4. Ghi nhận", 1)
    doc.add_paragraph("Mỗi khiếu nại có mã KN-2026-xxxxxx. Ghi lại nguyên văn yêu cầu của khách, cách giải quyết và lý do. "
                      "Khiếu nại khẩn cấp về Lock L9 được chuyển thẳng cho kỹ thuật viên trực, số nội bộ 8009.")
    doc.add_heading("5. Chỉ số theo dõi", 1)
    bullets(doc, [f"Mức hài lòng (CSAT) mục tiêu: {c.KPI['csat_target']}.",
                  f"Tỷ lệ giải quyết ngay cuộc gọi đầu: {c.KPI['first_call_resolution']}.",
                  f"Trả lời {c.KPI['answer_within']}; thời gian giữ máy tối đa {c.KPI['max_hold']}."])
    save_doc(doc, "noi-bo/Quy-trinh-xu-ly-khieu-nai.docx")


def call_scripts():
    doc = new_doc("Kịch bản trả lời tổng đài", "Mẫu câu cho các tình huống thường gặp")
    scripts = [
        ("Chào hỏi", f"Dạ, {c.COMPANY['name']} xin nghe, em là [tên]. Em có thể giúp gì cho anh/chị ạ?"),
        ("Khách hỏi đổi trả", f"Dạ, mình được đổi trả trong {c.RETURNS['days']} ngày kể từ ngày nhận hàng, hội viên Family và Pro là "
                              f"{c.RETURNS['member_days']} ngày. Anh/chị cho em xin mã đơn hàng để em kiểm tra ạ."),
        ("Khách báo khoá không mở", "Dạ, anh/chị đang đứng ngoài cửa phải không ạ? Mình thử cấp nguồn bằng sạc dự phòng qua cổng USB-C ở mặt dưới khoá, "
                                   "rồi nhập mã số. Nếu vẫn không được, em chuyển ngay cho kỹ thuật viên trực."),
        ("Khách hỏi tiến độ bảo hành", "Dạ, anh/chị đọc giúp em mã yêu cầu bắt đầu bằng BH. Thời gian xử lý là 7 ngày làm việc tại trung tâm ạ."),
        ("Khách muốn huỷ gói Cloud", "Dạ, mình huỷ được ngay trong ứng dụng, mục Tài khoản > Gói Cloud. Gói vẫn dùng được đến hết kỳ đã thanh toán ạ."),
        ("Khách bực bội", "Dạ, em rất tiếc về trải nghiệm này. Em ghi nhận và sẽ theo đến khi xong việc; em gọi lại cho anh/chị trong 2 giờ làm việc ạ."),
        ("Kết thúc", "Dạ, anh/chị còn cần em hỗ trợ gì nữa không ạ? Cảm ơn anh/chị đã gọi Mộc Lan Home."),
    ]
    for title, text in scripts:
        doc.add_heading(title, 2)
        doc.add_paragraph(text)
    doc.add_heading("Những điều không nói", 1)
    bullets(doc, ["Không hứa ngày giao linh kiện khi kho chưa xác nhận.", "Không hỏi mật khẩu tài khoản hay mã quản trị của khoá.",
                  "Không đọc số thẻ thanh toán của khách thành tiếng."])
    save_doc(doc, "noi-bo/Kich-ban-tong-dai.docx")


# --- Excel ----------------------------------------------------------------

def sheet(wb, title, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(list(row))
    for i, _ in enumerate(header, start=1):
        width = max(len(str(ws.cell(row=r, column=i).value or "")) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 60)
    ws.freeze_panes = "A2"
    return ws


def save_book(wb, path):
    target = OUT / path
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)


def price_list():
    wb = Workbook()
    ws = sheet(wb, "Sản phẩm", ["Mã", "Tên sản phẩm", "Loại", "Giá bán lẻ (đ)", "Giá đại lý cấp 1 (đ)", "Giá đại lý cấp 2 (đ)", "Bảo hành (tháng)"],
               [(p["code"], p["name"], p["kind"], p["price"], round(p["price"] * 0.78, -3), round(p["price"] * 0.85, -3), p["warranty"]) for p in c.PRODUCTS], first=True)
    ws.append(["", "Tổng giá trị một bộ đủ 8 sản phẩm", "", f"=SUM(D2:D{len(c.PRODUCTS) + 1})"])
    sheet(wb, "Phụ kiện", ["Mã", "Tên phụ kiện", "Giá bán lẻ (đ)", "Bảo hành (tháng)"], c.ACCESSORIES)
    sheet(wb, "Gói Cloud", ["Gói", "Giá tháng (đ)", "Giá năm (đ)", "Lưu video", "Số camera", "Phản hồi hỗ trợ"],
          [(n, p, p * 10, d, cams, sla) for n, p, d, cams, sla in c.PLANS])
    sheet(wb, "Khuyến mãi", ["Chương trình", "Từ ngày", "Đến ngày", "Nội dung", "Mã giảm giá"], c.PROMOTIONS)
    save_book(wb, "bang-gia/Bang-gia-2026.xlsx")


def spare_parts():
    parts = [("LK-C5-MOTOR", "C5", "Động cơ xoay Cam C5", 220_000, 90_000), ("LK-C5-LENS", "C5", "Cụm ống kính Cam C5", 310_000, 90_000),
             ("LK-O7-IR", "O7", "Bo đèn hồng ngoại Cam O7", 260_000, 120_000), ("LK-L9-FP", "L9", "Mặt đọc vân tay Lock L9", 890_000, 200_000),
             ("LK-L9-MOR", "L9", "Ruột khoá Lock L9", 650_000, 200_000), ("LK-L9-PIN", "L9", "Bộ pin và gioăng Lock L9", c.WARRANTY["battery_swap_l9"], 0),
             ("LK-M3-PSU", "M3", "Bộ nguồn Mesh M3", 150_000, 0), ("LK-D6-BAT", "D6", "Pin 6.500 mAh Bell D6", 380_000, 90_000),
             ("LK-H2-BAT", "H2", "Pin dự phòng Hub H2", 240_000, 90_000), ("LK-P1-RELAY", "P1", "Rơ-le Plug P1", 70_000, 50_000)]
    wb = Workbook()
    ws = sheet(wb, "Linh kiện", ["Mã linh kiện", "Sản phẩm", "Tên linh kiện", "Giá linh kiện (đ)", "Công thay (đ)", "Tổng (đ)"],
               [(*p, None) for p in parts], first=True)
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=6, value=f"=D{r}+E{r}")
    sheet(wb, "Ghi chú", ["Mục", "Nội dung"],
          [("Phí kiểm tra", f"{c.vnd(c.WARRANTY['out_of_warranty_check'])}, trừ vào phí sửa nếu khách đồng ý sửa"),
           ("Bảo hành linh kiện thay thế", "3 tháng kể từ ngày thay"),
           ("Sản phẩm ngừng cấp linh kiện", "Cam C3 và Hub H1 ngừng cấp linh kiện từ 31/12/2026")])
    save_book(wb, "bang-gia/Linh-kien-thay-the.xlsx")


def error_codes():
    wb = Workbook()
    sheet(wb, "Mã lỗi", ["Mã lỗi", "Sản phẩm", "Ý nghĩa", "Hướng dẫn khách", "Chuyển tiếp"],
          [(code, c.product(pc)["name"], meaning, action, level) for code, pc, meaning, action, level in c.ERROR_CODES], first=True)
    sheet(wb, "Đèn LED", ["Trạng thái đèn", "Ý nghĩa"], c.LED_STATES)
    sheet(wb, "Cấp hỗ trợ", ["Cấp", "Ai xử lý", "Phạm vi", "Thời hạn chuyển tiếp"], c.SUPPORT_LEVELS)
    sheet(wb, "Khôi phục gốc", ["Sản phẩm", "Cách khôi phục cài đặt gốc"], [(p["name"], p["reset"]) for p in c.PRODUCTS])
    save_book(wb, "ky-thuat/Ma-loi-va-den-bao.xlsx")


def install_fees():
    wb = Workbook()
    sheet(wb, "Phí lắp đặt", ["Sản phẩm", "Hạng mục", *[f"{z} (đ)" for z in c.INSTALL_ZONES]],
          [(c.product(pc)["name"], item, a, b, z) for pc, item, a, b, z in c.INSTALL_FEES], first=True)
    sheet(wb, "Điều kiện", ["Mục", "Nội dung"],
          [("Miễn phí", "Lắp Lock L9 lên cửa gỗ trong nội thành được miễn phí"),
           ("Đặt lịch", "Đặt trước ít nhất 24 giờ; kỹ thuật viên đến trong khung 2 giờ đã hẹn"),
           ("Huỷ lịch", "Huỷ trước 4 giờ không mất phí; huỷ muộn hơn tính 50.000 đ"),
           ("Ngoài giờ", "Sau 18:00 và Chủ nhật cộng 20% phí lắp"),
           ("Vật tư", "Dây điện, ống ghen và cáp mạng tính riêng theo mét: 12.000 đ/m")])
    save_book(wb, "dich-vu/Phi-lap-dat.xlsx")


def centres_and_dealers():
    wb = Workbook()
    sheet(wb, "Trung tâm bảo hành", ["Thành phố", "Địa chỉ", "Giờ làm việc", f"Tổng đài {c.COMPANY['hotline']}", "Dịch vụ"], c.CENTRES, first=True)
    sheet(wb, "Đại lý", ["Mã đại lý", "Tên đại lý", "Tỉnh thành", "Cấp"], c.DEALERS)
    save_book(wb, "dich-vu/Trung-tam-bao-hanh-va-dai-ly.xlsx")


def compatibility():
    wb = Workbook()
    sheet(wb, "Tương thích", ["Sản phẩm", *c.COMPATIBILITY["columns"]],
          [(c.product(code)["name"], *values) for code, values in c.COMPATIBILITY["rows"].items()], first=True)
    save_book(wb, "ky-thuat/Bang-tuong-thich.xlsx")


def firmware_history():
    wb = Workbook()
    sheet(wb, "Firmware", ["Sản phẩm", "Phiên bản", "Ngày phát hành", "Thay đổi"],
          [(c.product(pc)["name"], version, date, "; ".join(changes)) for version, pc, date, changes in c.FIRMWARE], first=True)
    save_book(wb, "ky-thuat/Lich-su-firmware.xlsx")


# --- Plain formats --------------------------------------------------------

def plain_files():
    folder = OUT / "khac"
    folder.mkdir(parents=True, exist_ok=True)
    with open(folder / "Danh-sach-dai-ly.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["ma_dai_ly", "ten_dai_ly", "tinh_thanh", "cap"])
        writer.writerows(c.DEALERS)
    (folder / "Nhat-ky-thay-doi-ung-dung.md").write_text(
        "# Nhật ký thay đổi ứng dụng Mộc Lan Home\n\n"
        "## 5.2.0 (2026-10-05)\n- Thêm chế độ Vắng nhà: một chạm bật mọi camera và khoá cửa.\n- Chia sẻ nhà cho tối đa 8 thành viên.\n\n"
        "## 5.1.3 (2026-09-12)\n- Sửa lỗi thông báo đến trễ trên Android 14.\n- Xem lại video nhanh gấp 4 lần.\n\n"
        "## 5.1.0 (2026-08-01)\n- Đăng nhập bằng khuôn mặt hoặc vân tay.\n- Widget bật tắt Plug P1 trên màn hình chính.\n\n"
        "## 5.0.0 (2026-06-18)\n- Giao diện mới. Yêu cầu iOS 15 hoặc Android 9 trở lên.\n", encoding="utf-8")
    (folder / "Thong-bao-bao-tri.txt").write_text(
        "THÔNG BÁO BẢO TRÌ HỆ THỐNG CLOUD\n\n"
        "Thời gian: 01:00 đến 03:30 ngày 18/10/2026.\n"
        "Ảnh hưởng: không xem lại được video trên Cloud và không nhận thông báo đẩy. "
        "Camera vẫn ghi vào thẻ nhớ; khoá cửa và kịch bản của hub hoạt động bình thường.\n"
        "Sau bảo trì, video ghi trong thời gian này được tải lên Cloud trong vòng 6 giờ.\n"
        f"Liên hệ: {c.COMPANY['hotline']} hoặc {c.COMPANY['email']}.\n", encoding="utf-8")
    payments = "".join(f"<li>{p}</li>" for p in c.PAYMENTS)
    (folder / "Cau-hoi-ve-thanh-toan.html").write_text(
        "<!doctype html><html lang=\"vi\"><head><meta charset=\"utf-8\"><title>Câu hỏi về thanh toán</title></head><body>"
        "<h1>Câu hỏi về thanh toán</h1>"
        f"<h2>Mộc Lan Home nhận những cách thanh toán nào?</h2><ul>{payments}</ul>"
        "<h2>Trả góp cần giấy tờ gì?</h2><p>Chỉ cần thẻ tín dụng của một trong 18 ngân hàng liên kết. Phí chuyển đổi trả góp là 0 đ.</p>"
        "<h2>Khi nào tôi nhận được hoá đơn điện tử?</h2><p>Hoá đơn được gửi qua email trong 24 giờ sau khi giao hàng. "
        "Muốn xuất hoá đơn công ty, hãy nhập mã số thuế lúc đặt hàng; không sửa được sau 7 ngày.</p>"
        "<h2>Tôi bị trừ tiền hai lần thì sao?</h2><p>Khoản trừ trùng được hoàn tự động trong 3 ngày làm việc. "
        "Nếu quá hạn, hãy gửi ảnh sao kê đến bộ phận hỗ trợ.</p></body></html>", encoding="utf-8")


def main():
    for build in (warranty_policy, cloud_terms, returns_en, lock_install_guide, hub_manual, complaint_process, call_scripts,
                  price_list, spare_parts, error_codes, install_fees, centres_and_dealers, compatibility, firmware_history, plain_files):
        build()
    files = sorted(p.relative_to(OUT) for p in OUT.rglob("*") if p.is_file())
    print(f"{len(files)} files in {OUT}")


if __name__ == "__main__":
    main()
